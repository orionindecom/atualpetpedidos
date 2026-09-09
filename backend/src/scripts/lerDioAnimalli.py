"""Leitura somente: requer Python 3 e openpyxl. Não recalcula nem altera o Excel."""
import json
import math
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
import openpyxl


def norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def extract(path):
    book = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    candidates = [s for s in book if any("tabela de valores dio animalli" in norm(c.value)
                  for row in s.iter_rows() for c in row)]
    if len(candidates) != 1:
        raise ValueError("Esperada uma única aba com título TABELA DE VALORES DIO ANIMALLI")
    sheet = candidates[0]
    result = {"aba": sheet.title, "abasIgnoradas": [s.title for s in book if s != sheet],
              "colunaPreco": "Valor Unitário", "registros": [], "ignoradas": []}
    header = None
    for row in sheet.iter_rows():
        values = [c.value for c in row]
        labels = [norm(v) for v in values]
        if "descrição do produto" in labels and "embal." in labels:
            result["ignoradas"].append({"linha": row[0].row, "motivo": "cabeçalho auxiliar"})
            continue
        if "valor unitário" in labels and "produto" in labels:
            if labels.count("valor unitário") != 1:
                raise ValueError("Cabeçalho de preço ambíguo")
            # A seção promocional repete PRODUTO em A/B; B continua sendo o nome.
            header = {"nome": max(i for i, v in enumerate(labels) if v == "produto"),
                      "valor": labels.index("valor unitário"),
                      "embalagem": labels.index("emb."), "diluicao": labels.index("diluição"),
                      "ean": labels.index("ean") if "ean" in labels else None}
            result["ignoradas"].append({"linha": row[0].row, "motivo": "cabeçalho"})
            continue
        base = {"linha": row[0].row}
        if not header:
            result["ignoradas"].append({**base, "motivo": "título/metadados antes do cabeçalho"})
            continue
        name, pack, price = [values[header[k]] for k in ("nome", "embalagem", "valor")]
        if name is None and pack is None and price is None or pack is None and price is None:
            result["ignoradas"].append({**base, "motivo": "linha vazia/título/metadados"})
            continue
        raw_ean = values[header["ean"]] if header["ean"] is not None else None
        ean = str(raw_ean).strip() if raw_ean is not None else ""
        record = {**base, "nome": str(name or "").strip(), "ean": ean,
                  "embalagem": str(pack or "").strip(),
                  "diluicao": str(values[header["diluicao"]] or "").strip(),
                  "celula": row[header["valor"]].coordinate, "valorOriginal": price}
        if not name or not pack or isinstance(price, bool) or not isinstance(price, (int, float)) or not math.isfinite(price) or price <= 0:
            result["ignoradas"].append({**record, "motivo": "produto/embalagem ausente ou Valor Unitário não numérico/positivo"})
            continue
        record["valor"] = float(Decimal(str(price)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        record["formula"] = formulas[sheet.title][record["celula"]].data_type == "f"
        result["registros"].append(record)
    if not header or not result["registros"]:
        raise ValueError("Nenhuma linha válida em Valor Unitário")
    return result


if __name__ == "__main__":
    print(json.dumps(extract(sys.argv[1]), ensure_ascii=True, allow_nan=False))
